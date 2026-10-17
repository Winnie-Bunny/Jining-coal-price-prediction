"""Convert customer Excel tables to validated model request records."""
import datetime as dt
from io import BytesIO
import math
from pathlib import Path
import re
import unicodedata
from zipfile import BadZipFile, ZipFile

ALIASES = {
    '京唐港瘦煤(加)（元/吨）': '京唐港瘦煤（加）（元/吨）',
    '龙煤鹤岗1/3焦煤(A8.5)（元/吨）': '龙煤集团鹤岗1/3焦煤(A8.5）（元/吨）',
    '龙煤鹤岗1/3焦煤(A9.5)（元/吨）': '龙煤集团鹤岗1/3焦煤（A9.5）（元/吨）',
    '上海港主焦煤(澳)（元/吨）': '上海港主焦煤（澳）（元/吨）',
    '天津港主焦煤(蒙古)（元/吨）': '天津港主焦煤（蒙古）（元/吨）',
    '天峻焦煤(G70)（元/吨）': '天峻焦煤（G70）（元/吨）',
    '天峻焦煤(G80)（元/吨）': '天峻焦煤（G80）（元/吨）',
    '乌海肥精煤(G>88)（元/吨）': '乌海肥精煤（G>88)（元/吨）',
    '乌海肥精煤(G90)（元/吨）': '乌海肥精煤（G90）（元/吨）',
    '冶金煤:内蒙古策克1/3焦煤（元/吨）': '冶金煤:内蒙古策克蒙古1/3焦煤（元/吨）',
    '日照港主焦煤(澳)（元/吨）': '日照港主焦煤（澳）（元/吨）',
}
COUNTS = {7: 60, 30: 90, 365: 364}


class ExcelImportError(ValueError):
    pass


def normalized(value):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', str(value or '')).lstrip('\ufeff')).casefold()


def columns_for(headers, feature_map):
    names = {'date': ('date', '日期', '时间'), 'history_price': ('history_price', '历史价格', '历史价格（元/吨）')}
    for code, chinese in feature_map.items():
        names[code] = (code, chinese, ALIASES.get(chinese, chinese))
    lookup = {normalized(alias): code for code, aliases in names.items() for alias in aliases}
    columns = {}
    for index, header in enumerate(headers):
        code = lookup.get(normalized(header))
        if code is not None:
            if code in columns:
                raise ExcelImportError(f'第 1 行有重复列：{header}。同一字段只保留一列。')
            columns[code] = index
    missing = [code for code in names if code not in columns]
    if missing:
        labels = [feature_map.get(code, {'date': '日期', 'history_price': '历史价格'}.get(code, code)) for code in missing]
        detail = '、'.join(labels[:8]) + ('…' if len(labels) > 8 else '')
        raise ExcelImportError(f'缺少 {len(missing)} 个必需列：{detail}。请使用下载的模板，第一行为列名。')
    return columns


def date_value(value, row):
    if isinstance(value, dt.datetime):
        if value.time() != dt.time():
            raise ExcelImportError(f'第 {row} 行日期包含具体时刻，请按天填写日期。')
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        for pattern in ('%Y-%m-%d', '%Y/%m/%d', '%Y年%m月%d日'):
            try:
                return dt.datetime.strptime(value, pattern).date().isoformat()
            except ValueError:
                pass
    raise ExcelImportError(f'第 {row} 行日期无效，请使用 Excel 日期或 YYYY-MM-DD。')


def numeric_value(value, row, label):
    if isinstance(value, str):
        value = value.strip()
        if re.fullmatch(r'[+-]?\d{1,3}(,\d{3})+(\.\d+)?', value):
            value = value.replace(',', '')
    try:
        if isinstance(value, bool) or value is None or value == '':
            raise ValueError()
        result = float(value)
        if not math.isfinite(result):
            raise ValueError()
        return result
    except (TypeError, ValueError, OverflowError):
        raise ExcelImportError(f'第 {row} 行「{label}」缺失或不是有效数字。公式请先在 Excel 中计算并保存，或粘贴为数值。') from None


def rows_to_request(rows, model_id, feature_map):
    if model_id not in COUNTS or isinstance(model_id, bool):
        raise ExcelImportError('预测周期必须为 7、30 或 365。')
    rows = iter(rows)
    try:
        headers = next(rows)
    except StopIteration:
        raise ExcelImportError('Excel 工作表为空。') from None
    columns = columns_for(headers, feature_map)
    by_date = {}
    duplicates = 0
    for row_number, values in enumerate(rows, 2):
        if row_number > 5001:
            raise ExcelImportError('Excel 最多支持 5000 行数据，请删除多余的空白格式或拆分文件。')
        if not any(value is not None and value != '' for value in values):
            continue
        record = {}
        for code, index in columns.items():
            value = values[index] if index < len(values) else None
            record[code] = date_value(value, row_number) if code == 'date' else numeric_value(value, row_number, feature_map.get(code, '历史价格'))
        date = record['date']
        if date in by_date:
            if record != by_date[date]:
                raise ExcelImportError(f'日期 {date} 存在不同数据，请确认后只保留一行。')
            duplicates += 1
        else:
            by_date[date] = record
    ordered = sorted(by_date.values(), key=lambda row: row['date'])
    for previous, current in zip(ordered, ordered[1:]):
        if (dt.date.fromisoformat(current['date']) - dt.date.fromisoformat(previous['date'])).days != 1:
            raise ExcelImportError(f'日期不连续：{previous["date"]} 与 {current["date"]} 之间缺少数据。')
    count = COUNTS[model_id]
    if len(ordered) < count:
        raise ExcelImportError(f'当前只有 {len(ordered)} 天有效数据，所选预测周期需要至少 {count} 天。')
    selected = ordered[-count:]
    return {'id': model_id, 'price_data': selected}, {
        'total_days': len(ordered), 'used_days': count, 'duplicates_removed': duplicates,
        'start_date': selected[0]['date'], 'end_date': selected[-1]['date'],
    }


def convert_excel(content, filename, model_id, feature_map):
    suffix = Path(filename or '').suffix.lower()
    if suffix not in ('.xlsx', '.xls'):
        raise ExcelImportError('仅支持 .xlsx 和 .xls 文件；请将其他格式另存为 .xlsx。')
    if not content or len(content) > 8 * 1024 * 1024:
        raise ExcelImportError('Excel 文件为空或超过 8 MB。')
    book = None
    try:
        if suffix == '.xlsx':
            from openpyxl import load_workbook
            with ZipFile(BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 32 * 1024 * 1024:
                    raise ExcelImportError('Excel 展开后过大，请删除多余工作表、图片和空白格式。')
            book = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
            sheet = book['数据'] if '数据' in book.sheetnames else book.worksheets[0]
            if (sheet.max_column or 0) > 512 or (sheet.max_row or 0) > 5001:
                raise ExcelImportError('数据区域过大：最多 5000 行数据、512 列。')
            rows = sheet.iter_rows(values_only=True)
        else:
            import xlrd
            book = xlrd.open_workbook(file_contents=content, on_demand=True)
            sheet = book.sheet_by_name('数据') if '数据' in book.sheet_names() else book.sheet_by_index(0)
            if sheet.nrows > 5001 or sheet.ncols > 512:
                raise ExcelImportError('数据区域过大：最多 5000 行数据、512 列。')
            rows = ([xlrd.xldate.xldate_as_datetime(cell.value, book.datemode) if cell.ctype == xlrd.XL_CELL_DATE else bool(cell.value) if cell.ctype == xlrd.XL_CELL_BOOLEAN else cell.value for cell in sheet.row(i)] for i in range(sheet.nrows))
        payload, summary = rows_to_request(rows, model_id, feature_map)
        summary['sheet'] = sheet.title if suffix == '.xlsx' else sheet.name
        return payload, summary
    except ExcelImportError:
        raise
    except Exception:
        raise ExcelImportError('无法读取 Excel，请确认文件未加密且没有损坏，并另存为 .xlsx 后重试。') from None
    finally:
        if book is not None:
            book.close() if suffix == '.xlsx' else book.release_resources()
