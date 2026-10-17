"""Exercise Excel normalization, real XLSX/XLS readers and input rejection."""
import datetime as dt
from io import BytesIO
import json
from pathlib import Path
import sys
import unittest
from zipfile import ZipFile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'inference'))
from excel_import import ExcelImportError, convert_excel, rows_to_request

MATERIALS = ROOT / 'project_materials_2026-10-16'
MAPPING = json.loads((MATERIALS / 'docs/feature-map.json').read_text())
MINI = {'FACTOR': '测试因素（元/吨）'}


def history(count=60):
    return [['日期', '历史价格', '测试因素（元/吨）']] + [
        [(dt.date(2025, 1, 1) + dt.timedelta(days=i)).isoformat(), 1000, 12] for i in range(count)]


class ExcelTests(unittest.TestCase):
    def test_real_xlsx_preserves_all_features_and_dates(self):
        payload, summary = convert_excel((MATERIALS / 'tests/fixtures/week.xlsx').read_bytes(), '客户数据.xlsx', 7, MAPPING)
        expected = json.loads((MATERIALS / 'examples/request-week.json').read_text())
        self.assertEqual(payload, expected)
        self.assertEqual(summary['used_days'], 60)
        self.assertEqual(summary['sheet'], '数据')

    def test_english_headers_sorting_duplicates_and_recent_horizon(self):
        rows = history(100)
        rows[0] = ['date', 'history_price', 'FACTOR']
        rows[1][1] = '1,000.25'
        rows = [rows[0], *reversed(rows[1:]), rows[-1]]
        payload, summary = rows_to_request(rows, 30, MINI)
        self.assertEqual(len(payload['price_data']), 90)
        self.assertEqual(payload['price_data'][-1]['date'], '2025-04-10')
        self.assertEqual(summary['duplicates_removed'], 1)
        annual, _ = rows_to_request(history(400), 365, MINI)
        self.assertEqual(len(annual['price_data']), 364)

    def test_missing_duplicate_columns_invalid_values_dates_and_gaps(self):
        cases = []
        rows = history(); rows[0][2] = '错误列名'; cases.append(rows)
        rows = history(); rows[0].append('FACTOR'); cases.append(rows)
        for value in (None, '', True, 'NaN', '文本', '88%'):
            rows = history(); rows[1][2] = value; cases.append(rows)
        rows = history(); rows[1][0] = 'bad'; cases.append(rows)
        rows = history(); rows.pop(20); cases.append(rows)
        rows = history(); rows.append(['2025-01-01', 1001, 12]); cases.append(rows)
        cases.append(history(59))
        for rows in cases:
            with self.subTest(first=rows[1]):
                with self.assertRaises(ExcelImportError):
                    rows_to_request(rows, 7, MINI)

    def test_real_legacy_xls_dates(self):
        # Artifact Tool exports XLSX only; xlwt creates this disposable BIFF fixture.
        import xlwt
        book = xlwt.Workbook(); sheet = book.add_sheet('数据')
        style = xlwt.easyxf(num_format_str='YYYY-MM-DD')
        for row_idx, row in enumerate(history()):
            for col_idx, value in enumerate(row):
                if row_idx and col_idx == 0:
                    sheet.write(row_idx, col_idx, dt.datetime.fromisoformat(value), style)
                else:
                    sheet.write(row_idx, col_idx, value)
        stream = BytesIO(); book.save(stream)
        payload, _ = convert_excel(stream.getvalue(), '旧数据.xls', 7, MINI)
        self.assertEqual(payload['price_data'][0]['date'], '2025-01-01')
        self.assertEqual(payload['price_data'][0]['FACTOR'], 12)

    def test_corrupt_files_and_uncalculated_formula_are_rejected(self):
        for filename in ('bad.xlsx', 'bad.xls', 'bad.csv'):
            with self.subTest(filename=filename), self.assertRaises(ExcelImportError):
                convert_excel(b'not a spreadsheet', filename, 7, MINI)
        ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        output = BytesIO()
        with ZipFile(MATERIALS / 'tests/fixtures/week.xlsx') as original, ZipFile(output, 'w') as modified:
            for name in original.namelist():
                content = original.read(name)
                if name == 'xl/worksheets/sheet1.xml':
                    tree = ET.fromstring(content)
                    cell = tree.find(".//m:c[@r='B2']", ns)
                    for child in list(cell): cell.remove(child)
                    ET.SubElement(cell, '{'+ns['m']+'}f').text = '1+1'
                    content = ET.tostring(tree)
                modified.writestr(name, content)
        with self.assertRaisesRegex(ExcelImportError, '第 2 行'):
            convert_excel(output.getvalue(), 'formula.xlsx', 7, MAPPING)

    def test_blank_template_has_complete_headers_and_no_fake_prices(self):
        import openpyxl
        path = ROOT / 'inference/static/downloads/coal-price-template.xlsx'
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = book['数据']
            rows = sheet.iter_rows(values_only=True)
            self.assertEqual(len(next(rows)), 213)
            self.assertTrue(all(all(value is None for value in row) for row in rows))
        finally:
            book.close()


if __name__ == '__main__':
    unittest.main()
